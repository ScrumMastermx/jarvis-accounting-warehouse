#!/usr/bin/env python3
"""Conciliación bancaria: estado de cuenta (banco) vs libro/auxiliar contable.

Uso: python scripts/conciliar.py <banco.csv> <libro.csv> <salida.xlsx> [--dias N]

banco.csv: fecha,descripcion,monto     libro.csv: fecha,referencia,monto
Fechas en AAAA-MM-DD (también acepta DD/MM/AAAA). Montos con signo: cargos negativos, abonos positivos.
Empareja por monto exacto (2 decimales) y fecha dentro de ±N días (3 por defecto). Cada movimiento se usa una vez.
"""
import csv
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


def parse_fecha(s: str) -> datetime:
    s = s.strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    raise ValueError(f"fecha no reconocida: {s!r} (usa AAAA-MM-DD)")


def parse_monto(s: str) -> Decimal:
    s = s.strip().replace("$", "").replace(",", "")
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except InvalidOperation:
        raise ValueError(f"monto no reconocido: {s!r}")


def leer(path: Path, col_texto: str):
    filas = []
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rd = csv.DictReader(fh)
        faltan = {"fecha", col_texto, "monto"} - set(h.strip().lower() for h in rd.fieldnames or [])
        if faltan:
            raise SystemExit(f"ERROR: a {path} le faltan columnas: {', '.join(sorted(faltan))}")
        for n, row in enumerate(rd, start=2):
            row = {k.strip().lower(): (v or "").strip() for k, v in row.items() if k}
            if not any(row.values()):
                continue
            filas.append({"fila": n, "fecha": parse_fecha(row["fecha"]), "texto": row[col_texto],
                          "monto": parse_monto(row["monto"]), "usado": False})
    return filas


def escribir(ws, encabezados, filas):
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in filas:
        ws.append(f)
    ws.freeze_panes = "A2"
    for col, ancho in zip("ABCDEFG", (12, 40, 14, 8, 12, 40, 14)):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            if isinstance(c, type(fila[0])) and isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                c.number_format = "#,##0.00"
            if isinstance(c.value, datetime):
                c.number_format = "yyyy-mm-dd"


def main(argv):
    args = [a for a in argv[1:] if not a.startswith("--")]
    dias = 3
    if "--dias" in argv:
        dias = int(argv[argv.index("--dias") + 1])
        args = [a for a in args if a != str(dias)]
    if len(args) != 3:
        print(__doc__, file=sys.stderr)
        return 2
    banco = leer(Path(args[0]), "descripcion")
    libro = leer(Path(args[1]), "referencia")
    salida = Path(args[2])

    conciliadas = []
    for b in banco:
        candidatos = [l for l in libro if not l["usado"] and l["monto"] == b["monto"] and abs((l["fecha"] - b["fecha"]).days) <= dias]
        if candidatos:
            l = min(candidatos, key=lambda l: abs((l["fecha"] - b["fecha"]).days))
            l["usado"] = True
            b["usado"] = True
            conciliadas.append([b["fecha"], b["texto"], float(b["monto"]), b["fila"], l["fecha"], l["texto"], l["fila"]])
    solo_banco = [[b["fecha"], b["texto"], float(b["monto"]), b["fila"]] for b in banco if not b["usado"]]
    solo_libro = [[l["fecha"], l["texto"], float(l["monto"]), l["fila"]] for l in libro if not l["usado"]]

    wb = Workbook()
    ws = wb.active
    ws.title = "Conciliadas"
    escribir(ws, ["Fecha banco", "Descripción banco", "Monto", "Fila banco", "Fecha libro", "Referencia libro", "Fila libro"], conciliadas)
    escribir(wb.create_sheet("Solo banco"), ["Fecha", "Descripción", "Monto", "Fila"], solo_banco)
    escribir(wb.create_sheet("Solo libro"), ["Fecha", "Referencia", "Monto", "Fila"], solo_libro)

    rs = wb.create_sheet("Resumen")
    rs.append(["Concepto", "Movimientos", "Suma"])
    for c in rs[1]:
        c.font = Font(bold=True)
    rs.append(["Banco (total)", len(banco), float(sum(b["monto"] for b in banco))])
    rs.append(["Libro (total)", len(libro), float(sum(l["monto"] for l in libro))])
    rs.append(["Conciliadas", len(conciliadas), float(sum(Decimal(str(c[2])) for c in conciliadas))])
    rs.append(["Solo banco", len(solo_banco), float(sum(Decimal(str(c[2])) for c in solo_banco))])
    rs.append(["Solo libro", len(solo_libro), float(sum(Decimal(str(c[2])) for c in solo_libro))])
    rs.append(["Diferencia banco - libro", "", "=C2-C3"])
    rs.append(["Comprobación: solo banco - solo libro", "", "=C5-C6"])
    rs.append([])
    rs.append(["Criterio", f"monto exacto y fecha ±{dias} días"])
    rs.append(["Archivos", f"{args[0]} | {args[1]}"])
    rs.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M")])
    rs.column_dimensions["A"].width = 40
    rs.column_dimensions["B"].width = 14
    rs.column_dimensions["C"].width = 16
    for fila in rs.iter_rows(min_row=2, max_row=8):
        fila[2].number_format = "#,##0.00"

    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    print(f"{len(conciliadas)} conciliadas, {len(solo_banco)} solo banco, {len(solo_libro)} solo libro -> {salida}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
