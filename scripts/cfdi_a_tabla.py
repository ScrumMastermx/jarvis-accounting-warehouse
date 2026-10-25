#!/usr/bin/env python3
"""Convierte una carpeta de CFDI 4.0 (XML del SAT) en un Excel con hojas Facturas, Resumen y Fuente.

Uso: python scripts/cfdi_a_tabla.py <carpeta_con_xml> <salida.xlsx>

Los XML que no se puedan leer o no sean CFDI 4.0 se reportan en pantalla (stderr) y no detienen el proceso.
Los UUID repetidos se marcan en la columna "Duplicado" y NO se cuentan en el Resumen.
"""
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

NS_CFDI = "http://www.sat.gob.mx/cfd/4"
NS_TFD = "http://www.sat.gob.mx/TimbreFiscalDigital"

COLUMNAS = [
    "UUID", "Fecha", "Tipo", "Serie", "Folio", "Emisor RFC", "Emisor nombre", "Receptor RFC",
    "Receptor nombre", "Uso CFDI", "Subtotal", "Descuento", "IVA trasladado", "Total", "Moneda",
    "Forma de pago", "Método de pago", "Duplicado", "Archivo",
]


def leer_cfdi(path: Path) -> dict:
    root = ET.parse(path).getroot()
    if root.tag != f"{{{NS_CFDI}}}Comprobante" or root.get("Version") != "4.0":
        raise ValueError("no es un cfdi:Comprobante versión 4.0")
    emisor = root.find(f"{{{NS_CFDI}}}Emisor")
    receptor = root.find(f"{{{NS_CFDI}}}Receptor")
    timbre = root.find(f".//{{{NS_TFD}}}TimbreFiscalDigital")
    if emisor is None or receptor is None or timbre is None:
        raise ValueError("falta Emisor, Receptor o TimbreFiscalDigital")

    iva = 0.0
    for tras in root.iterfind(f"{{{NS_CFDI}}}Impuestos/{{{NS_CFDI}}}Traslados/{{{NS_CFDI}}}Traslado"):
        if tras.get("Impuesto") == "002":
            iva += float(tras.get("Importe", "0") or 0)

    fecha = datetime.fromisoformat(root.get("Fecha"))
    return {
        "UUID": timbre.get("UUID", "").upper(),
        "Fecha": fecha,
        "Tipo": root.get("TipoDeComprobante", ""),
        "Serie": root.get("Serie", ""),
        "Folio": root.get("Folio", ""),
        "Emisor RFC": emisor.get("Rfc", ""),
        "Emisor nombre": emisor.get("Nombre", ""),
        "Receptor RFC": receptor.get("Rfc", ""),
        "Receptor nombre": receptor.get("Nombre", ""),
        "Uso CFDI": receptor.get("UsoCFDI", ""),
        "Subtotal": float(root.get("SubTotal", "0") or 0),
        "Descuento": float(root.get("Descuento", "0") or 0),
        "IVA trasladado": iva,
        "Total": float(root.get("Total", "0") or 0),
        "Moneda": root.get("Moneda", ""),
        "Forma de pago": root.get("FormaPago", ""),
        "Método de pago": root.get("MetodoPago", ""),
        "Duplicado": "",
        "Archivo": path.name,
    }


def escribir_tabla(ws, encabezados, filas):
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True)
    for fila in filas:
        ws.append(fila)
    ws.freeze_panes = "A2"
    for i, h in enumerate(encabezados, 1):
        ancho = max([len(str(h))] + [len(str(r[i - 1])) for r in filas if r[i - 1] is not None][:200])
        ws.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 10), 50)


def main(argv):
    if len(argv) != 3:
        print(__doc__, file=sys.stderr)
        return 2
    carpeta, salida = Path(argv[1]), Path(argv[2])
    if not carpeta.is_dir():
        print(f"ERROR: {carpeta} no es una carpeta", file=sys.stderr)
        return 2

    archivos = sorted(p for p in carpeta.rglob("*") if p.suffix.lower() == ".xml")
    facturas, invalidos = [], []
    for p in archivos:
        try:
            facturas.append(leer_cfdi(p))
        except (ET.ParseError, ValueError, TypeError) as e:
            invalidos.append((p.name, str(e)))
            print(f"XML inválido: {p.name}: {e}", file=sys.stderr)

    vistos = set()
    duplicados = 0
    for f in facturas:
        if f["UUID"] in vistos:
            f["Duplicado"] = "SÍ"
            duplicados += 1
        else:
            vistos.add(f["UUID"])
    facturas.sort(key=lambda f: (f["Fecha"], f["UUID"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"
    escribir_tabla(ws, COLUMNAS, [[f[c] for c in COLUMNAS] for f in facturas])
    for col in ("Subtotal", "Descuento", "IVA trasladado", "Total"):
        letra = get_column_letter(COLUMNAS.index(col) + 1)
        for celda in ws[letra][1:]:
            celda.number_format = "#,##0.00"
    for celda in ws[get_column_letter(COLUMNAS.index("Fecha") + 1)][1:]:
        celda.number_format = "yyyy-mm-dd hh:mm"

    # Resumen: solo facturas no duplicadas
    por_mes = defaultdict(lambda: [0, 0.0, 0.0, 0.0])
    por_emisor = defaultdict(lambda: [0, 0.0, 0.0, 0.0])
    for f in facturas:
        if f["Duplicado"]:
            continue
        for key, bucket in ((f["Fecha"].strftime("%Y-%m"), por_mes), ((f["Emisor RFC"], f["Emisor nombre"]), por_emisor)):
            b = bucket[key]
            b[0] += 1
            b[1] += f["Subtotal"]
            b[2] += f["IVA trasladado"]
            b[3] += f["Total"]

    rs = wb.create_sheet("Resumen")
    rs.append(["Por mes", "Facturas", "Subtotal", "IVA", "Total"])
    for c in rs[1]:
        c.font = Font(bold=True)
    inicio = rs.max_row + 1
    for mes in sorted(por_mes):
        rs.append([mes] + por_mes[mes])
    fin = rs.max_row
    if fin >= inicio:
        rs.append(["TOTAL", f"=SUM(B{inicio}:B{fin})", f"=SUM(C{inicio}:C{fin})", f"=SUM(D{inicio}:D{fin})", f"=SUM(E{inicio}:E{fin})"])
        for c in rs[rs.max_row]:
            c.font = Font(bold=True)
    rs.append([])
    rs.append(["Emisor RFC", "Emisor nombre", "Facturas", "Subtotal", "IVA", "Total"])
    for c in rs[rs.max_row]:
        c.font = Font(bold=True)
    inicio = rs.max_row + 1
    for (rfc, nombre), v in sorted(por_emisor.items(), key=lambda kv: -kv[1][3]):
        rs.append([rfc, nombre] + v)
    fin = rs.max_row
    if fin >= inicio:
        rs.append(["TOTAL", "", f"=SUM(C{inicio}:C{fin})", f"=SUM(D{inicio}:D{fin})", f"=SUM(E{inicio}:E{fin})", f"=SUM(F{inicio}:F{fin})"])
        for c in rs[rs.max_row]:
            c.font = Font(bold=True)
    for col in "BCDEF":
        rs.column_dimensions[col].width = 16
    rs.column_dimensions["A"].width = 18
    for fila in rs.iter_rows(min_row=2):
        for c in fila:
            if isinstance(c.value, float) or (isinstance(c.value, str) and c.value.startswith("=SUM")):
                c.number_format = "#,##0.00"

    fu = wb.create_sheet("Fuente")
    fu.append(["Carpeta", str(carpeta)])
    fu.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M")])
    fu.append(["XML encontrados", len(archivos)])
    fu.append(["Facturas leídas", len(facturas)])
    fu.append(["Duplicados (UUID repetido, fuera del Resumen)", duplicados])
    fu.append(["XML inválidos", len(invalidos)])
    for nombre, motivo in invalidos:
        fu.append(["  " + nombre, motivo])
    fu.column_dimensions["A"].width = 48
    fu.column_dimensions["B"].width = 60

    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    print(f"{len(facturas)} facturas ({duplicados} duplicadas, {len(invalidos)} XML inválidos) -> {salida}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
